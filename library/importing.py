"""Bounded, create-only import. Preview is stored server-side in Django session."""
import csv
import io
import hashlib
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from django.db import transaction
from .forms import StudentForm
from .models import Student, ImportBatch
from . import services

HEADERS=['full_name','grade','group','address','parent_name','parent_phone']
MAX_ROWS=500

def parse_upload(file=None,pasted=''):
    raw=file.read(2*1024*1024+1) if file else pasted.encode('utf-8')
    if len(raw)>2*1024*1024:raise services.DomainError('Ҳаҷм аз 2 MB зиёд аст.')
    if file and file.name.lower().endswith('.xlsx'):
        try:
            with ZipFile(io.BytesIO(raw)) as z:
                if len(z.infolist())>100 or sum(i.file_size for i in z.infolist())>10*1024*1024:
                    raise services.DomainError('Файли Excel аз ҳад калон аст.')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
            sheet=wb.active
            rows=[]
            if (sheet.max_row or 0)>MAX_ROWS+1:raise services.DomainError('Дар Excel то 500 сатр иҷозат аст; сатрҳои холии зиёдатиро тоза кунед.')
            for row in sheet.iter_rows(max_col=9,values_only=True):
                if len(rows)>MAX_ROWS:raise services.DomainError('Дар як импорт то 500 хонанда иҷозат аст.')
                row=list(row)
                while row and row[-1] is None:row.pop()
                rows.append(row)
            wb.close()
        except (BadZipFile,ValueError,KeyError,OSError,InvalidFileException,ParseError) as exc:raise services.DomainError('Файли Excel хонда нашуд.') from exc
    else:
        if file and not file.name.lower().endswith('.csv'):raise services.DomainError('Формати CSV ё XLSX лозим аст.')
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:raise services.DomainError('CSV бояд UTF-8 бошад.') from exc
        first=text.splitlines()[0] if text.strip() else ''
        delimiter='\t' if '\t' in first else (';' if first.count(';')>first.count(',') else ',')
        try:rows=list(csv.reader(io.StringIO(text),delimiter=delimiter,strict=True))
        except csv.Error as exc:raise services.DomainError('CSV нодуруст аст ё як майдон аз ҳад калон мебошад.') from exc
    if not rows:raise services.DomainError('Рӯйхат холӣ аст.')
    headers=[str(v or '').strip().lower() for v in rows[0]]
    allowed=set(HEADERS)|{'language','code'}
    if len(set(headers))!=len(headers) or not {'full_name','grade','group','address'}.issubset(headers) or not set(headers).issubset(allowed):
        raise services.DomainError('Сутунҳои full_name, grade, group, address лозиманд; parent_name ва parent_phone ихтиёрӣ. Намунаро боргирӣ кунед.')
    body=[r for r in rows[1:] if any(v is not None and str(v).strip() for v in r)]
    if any(len(row)>len(headers) for row in body):raise services.DomainError('Дар баъзе сатрҳо сутуни зиёдатӣ ҳаст.')
    if not body or len(body)>MAX_ROWS:raise services.DomainError('Рӯйхат бояд 1–500 хонанда дошта бошад.')
    return [{h:str(r[i] if i<len(r) and r[i] is not None else '').strip() for i,h in enumerate(headers)} for r in body],hashlib.sha256(raw).hexdigest()

def validate_rows(rows,school):
    seen=set();profiles=set();existing=set(Student.objects.filter(school=school,code__in=[r.get('code','') for r in rows]).values_list('code',flat=True));result=[];valid=[]
    def signature(data):
        return tuple(str(data.get(k,'')).strip().casefold() for k in ['full_name','grade','group','address','parent_phone'])
    from .forms import normalize_group
    existing_profiles=set()
    for student in Student.objects.filter(school=school).prefetch_related('enrollments'):
        for en in student.enrollments.all():
            if en.academic_year==school.academic_year:
                existing_profiles.add(signature(dict(full_name=student.full_name,grade=en.grade,group=normalize_group(en.group),address=student.address,parent_phone=student.parent_phone)))
    for number,row in enumerate(rows,2):
        form=StudentForm(row);ok=form.is_valid();errors=[]
        if not ok:errors=[f'{form.fields[k].label}: {", ".join(v)}' if k in form.fields else ', '.join(v) for k,v in form.errors.items()]
        code=row.get('code','')
        if code and (code in seen or code in existing):errors.append('Ин сатр аллакай ворид шудааст.')
        if code:seen.add(code)
        profile=signature(form.cleaned_data if ok else row)
        if profile in profiles:errors.append('Сатри хонанда дар файл такрор шудааст.')
        if profile in existing_profiles:errors.append('Хонанда бо ҳамин маълумот аллакай сабт шудааст.')
        profiles.add(profile)
        if ok and code:form.cleaned_data['code']=code
        result.append({'number':number,'row':row,'errors':errors})
        if not errors:valid.append(form.cleaned_data)
    return result,valid

@transaction.atomic
def commit_import(*,user,school,rows,token,file_hash):
    services.authorize(user,school,['admin','librarian']);services.lock_school(school)
    old=ImportBatch.objects.filter(token=token).first()
    if old:
        if old.school_id!=school.pk or old.file_hash!=file_hash:raise services.DomainError('Импорт ба амали дигар тааллуқ дорад.','conflict',409)
        return 0
    checked,valid=validate_rows(rows,school)
    if len(valid)!=len(checked):raise services.DomainError('Маълумот тағйир ёфт ё такрор ёфт шуд. Пешнамоиши нав лозим аст.')
    for row in valid:services.add_student(user=user,school=school,**row)
    batch=ImportBatch.objects.create(school=school,token=token,file_hash=file_hash,state='committed',created_by=user)
    services.audit(school,user,'students.imported',batch.pk,str(len(valid)))
    return len(valid)
